"""
generar_entrega.py - Arma el documento de entrega del equipo DSW-E01 (docs/entrega/P02_EQUIPO_01.docx)
a partir de las capturas y salidas que ya estan en docs/evidencia/. Requiere python-docx.
Uso: python scripts/generar_entrega.py
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "docs" / "evidencia" / "img"
TXT = ROOT / "docs" / "evidencia" / "txt"
OUT = ROOT / "docs" / "entrega" / "P02_EQUIPO_01.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

VERDE = RGBColor(0x2F, 0x6B, 0x3A)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---------- estilos base ----------
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
    h = doc.styles[f"Heading {nivel}"]
    h.font.name = "Calibri"
    h.font.size = Pt(tam)
    h.font.bold = True
    h.font.color.rgb = VERDE
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
    h.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.2)


def encabezado_pie():
    hp = sec.header.paragraphs[0]
    hp.text = "DSW-19559 · P02 · Equipo DSW-E01 · PR09 Monitoreo de huerto o ambiente"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = GRIS
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS


encabezado_pie()


# ---------- ayudas ----------
def p(texto, negrita=False, cursiva=False, tam=None, alinear=None, color=None, despues=None):
    par = doc.add_paragraph()
    r = par.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    if tam:
        r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color
    if alinear:
        par.alignment = alinear
    if despues is not None:
        par.paragraph_format.space_after = Pt(despues)
    return par


def h(texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


RECORTES = OUT.parent / ".recortes"


def recortar(ruta):
    """Quita el espacio vacio inferior de una captura de pagina completa (fondo uniforme)."""
    try:
        from PIL import Image
    except ImportError:
        return ruta
    im = Image.open(ruta).convert("RGB")
    w, h = im.size
    px = im.load()
    fondo = px[w - 1, h - 1]
    y = h - 1
    while y > 40:
        fila_vacia = all(abs(px[x, y][c] - fondo[c]) < 8 for x in range(0, w, 6) for c in range(3))
        if not fila_vacia:
            break
        y -= 1
    y = min(h, y + 28)
    if y >= h - 5:
        return ruta
    RECORTES.mkdir(exist_ok=True)
    salida = RECORTES / ruta.name
    im.crop((0, 0, w, y)).save(salida)
    return salida


def figura(nombre_png, pie, ancho_cm=16.5):
    ruta = IMG / nombre_png
    if not ruta.exists():
        p(f"[falta la imagen {nombre_png}]", cursiva=True, color=GRIS)
        return
    ruta = recortar(ruta)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
    figura.n += 1
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {figura.n}. {pie}")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GRIS
    c.paragraph_format.space_after = Pt(10)


figura.n = 0


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def tabla(encabezados, filas, anchos=None, tam=9.5):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(e)
        r.bold = True
        r.font.size = Pt(tam)
        sombrear(c, "E6F2E8")
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(tam)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bloque_codigo(texto, tam=8.5):
    lineas = texto.rstrip("\n").split("\n")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    sombrear(c, "F4F4F2")
    c.text = ""
    par = c.paragraphs[0]
    for i, ln in enumerate(lineas):
        r = par.add_run(ln)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(tam)
        if i < len(lineas) - 1:
            r.add_break()
    par.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def leer(nombre, max_lineas=None, desde=None, hasta=None):
    ruta = TXT / nombre
    if not ruta.exists():
        return f"[no se encontro {nombre}]"
    lineas = ruta.read_text(encoding="utf-8", errors="replace").splitlines()
    if desde is not None or hasta is not None:
        lineas = lineas[desde:hasta]
    if max_lineas:
        lineas = lineas[:max_lineas]
    return "\n".join(l.rstrip() for l in lineas)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- portada ----------
for _ in range(5):
    doc.add_paragraph()
p("Universidad Veracruzana", negrita=True, tam=14, alinear=WD_ALIGN_PARAGRAPH.CENTER, color=VERDE, despues=0)
p("Desarrollo de Sistemas Web (DSW-19559)", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=30)
p("P02. Aplicación Web 1.0 con JSP, Tomcat y PostgreSQL", negrita=True, tam=20, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=6)
p("Primer incremento del proyecto PR09", tam=14, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p("Monitoreo de huerto o ambiente", tam=14, cursiva=True, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=40)
p("Equipo DSW-E01", negrita=True, tam=13, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=4)
for nombre in ("Alejandro Pacheco Luna", "Juan Pablo Kuri Ricardez", "Ariadna Trejo Álvarez", "Pedro García Padilla"):
    p(nombre, tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
for _ in range(3):
    doc.add_paragraph()
p("Facilitador: Dr. Gabriel Rodríguez Vásquez", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p("Repositorio: github.com/zywoxxx/dsw-p02-monitoreo-huerto", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p("Xalapa, Ver., 11 de septiembre de 2026", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER)
salto()

# ---------- contenido ----------
h("Contenido")
for linea in (
    "1. Qué entregamos",
    "2. Cómo elegimos el proyecto (F02)",
    "3. Acta breve del proyecto",
    "4. Modelo de datos",
    "5. Cómo está construida la aplicación",
    "6. Instalación y despliegue",
    "7. Pruebas y evidencia",
    "8. Problemas que tuvimos y cómo los resolvimos",
    "9. Relación con la rúbrica R02",
    "10. Quién hizo qué",
    "11. Limitaciones y pendientes",
    "Anexo. Comandos y archivos de referencia",
):
    p(linea, despues=1)
salto()

# 1
h("1. Qué entregamos")
p("Este documento acompaña al repositorio del equipo y resume el primer incremento de nuestro proyecto "
  "integrador. El proyecto que nos tocó del banco C11 es el PR09, Monitoreo de huerto o ambiente. En esta "
  "primera etapa (Web 1.0) construimos una aplicación con JSP y Servlets que corre en Tomcat 9 y guarda todo en "
  "PostgreSQL 16. La idea es sencilla: alguien del huerto captura la lectura de un sensor, el sistema la valida, "
  "la guarda con su fecha y su procedencia, y si el valor se sale del rango que definimos para ese sensor, genera "
  "una alerta. Además se pueden dejar anotaciones por zona.")
p("Todo lo que aparece aquí se puede reproducir con el repositorio: el código, los scripts SQL, un archivo "
  "docker-compose para levantar la base de datos, un script que compila, despliega y corre las pruebas de "
  "extremo a extremo, y las capturas que se muestran más adelante. Los datos son ficticios. No hay contraseñas "
  "reales en ningún archivo versionado.")
tabla(["Elemento", "Dónde está"], [
    ("Código JSP/Servlet y configuración", "src/ y pom.xml (WAR web1.war)"),
    ("Modelo de datos y scripts SQL", "sql/01_schema.sql, 02_seed.sql, 03_consultas_verificacion.sql, 99_cleanup.sql"),
    ("README de instalación y despliegue", "README.md"),
    ("Acta breve, selección F02, trazabilidad, bitácora", "docs/"),
    ("Capturas y salidas de consola", "docs/evidencia/img/ y docs/evidencia/txt/"),
    ("Este documento", "docs/entrega/P02_EQUIPO_01.docx"),
], anchos=[6.5, 10])

# 2
h("2. Cómo elegimos el proyecto (F02)")
p("Antes de programar comparamos tres fichas del banco. Nos interesaba un problema que entendiéramos bien, "
  "que se pudiera simular sin hardware y que no nos obligara a resolver autorización por roles desde el primer "
  "incremento.")
p("El PR01 (inventario) lo descartamos porque su riesgo principal es la alteración de existencias y la ficha pide "
  "bitácora inmutable y permisos por rol; eso nos parecía mucho para arrancar. El PR03 (laboratorio de cómputo) "
  "nos gustó porque comparte el patrón de lecturas, umbrales y alertas, pero mezcla inventario de equipos e "
  "incidencias y su complejidad orientativa es media-alta. El PR09 tiene el mismo patrón, complejidad media, "
  "variables físicas fáciles de explicar (temperatura, humedad, luz) y un riesgo que se controla con una sola "
  "decisión de diseño: etiquetar cada lectura con su procedencia para no presentar una simulación como si fuera "
  "una medición real.")
p("Por eso pedimos el PR09 como opción principal y dejamos el PR03 como segunda opción. La publicación completa "
  "con la plantilla del foro está en docs/F02-seleccion-proyecto.md.")

# 3
h("3. Acta breve del proyecto")
tabla(["Campo", "Valor"], [
    ("Proyecto", "PR09 – Monitoreo de huerto o ambiente (C11)"),
    ("Equipo", "DSW-E01"),
    ("Actividad", "P02, del 31 de agosto al 11 de septiembre de 2026, rúbrica R02"),
    ("Stack fijo", "Java 11, Maven 3.9.9, Tomcat 9.0.115 (Servlet 4.0 / JSP 2.3), PostgreSQL 16, JDBC"),
    ("Roles funcionales", "Responsable de huerto, observador y coordinación (máximo tres, como pide la ficha)"),
], anchos=[4.5, 12])
h("Problema y usuarios", 2)
p("En un huerto escolar las observaciones se anotan en libretas o se mandan por mensaje: no queda claro cuándo se "
  "tomó cada dato, en qué unidad ni quién lo midió, y nadie se entera a tiempo cuando algo se sale de rango. Los "
  "usuarios son el responsable del huerto (define zonas, sensores y umbrales), el observador (captura lecturas y "
  "anotaciones) y coordinación (consulta el historial y las alertas).")
h("Alcance de este incremento", 2)
p("Incluimos zonas, variables con unidad y rango físico, sensores por zona, lecturas con instante y procedencia, "
  "umbral por sensor, historial, alertas automáticas y anotaciones por rol, más una ruta de salud para comprobar la "
  "conexión a la base. Dejamos fuera lo que la ficha excluye (riego automático, recomendaciones agronómicas, "
  "hardware obligatorio) y lo que corresponde a incrementos posteriores: autenticación, alta de zonas y sensores "
  "desde pantalla, atención de alertas, la API (RF08) y el simulador IoT (RF09).")
h("Flujo principal", 2)
p("El observador entra a /web1/lecturas y ve los sensores agrupados por zona con su variable y su umbral, y debajo "
  "el historial. Elige un sensor, escribe el valor y, si quiere, una observación. Al enviar, el servidor revisa que "
  "el sensor exista y esté activo, que el valor sea numérico con máximo dos decimales y que esté dentro del rango "
  "físico de la variable. Si algo falla, responde 400 y muestra la lista de errores sin perder lo que ya se había "
  "escrito. Si todo está bien, guarda la lectura marcada como manual; si el valor queda fuera del umbral del sensor "
  "crea también una alerta (BAJA o ALTA) en la misma transacción, y redirige con 303 a la lista, donde aparece el "
  "mensaje de confirmación y la fila nueva.")
h("Requisitos funcionales de la ficha y cómo los cubrimos", 2)
tabla(["RF", "Ficha PR09", "Qué hicimos en P02", "Cómo se comprueba"], [
    ("RF01", "Zonas", "Tabla zona; se listan en /anotaciones y aparecen en el catálogo de sensores", "Figuras 2 y 10"),
    ("RF02", "Cultivos o variables", "Tabla variable con unidad y rango físico; cada sensor muestra la suya", "Figura 2"),
    ("RF03", "Lecturas", "Formulario POST /lecturas con validación y procedencia manual", "Figuras 3, 4 y 7 a 9"),
    ("RF04", "Umbrales", "Tabla umbral (uno por sensor), visible en el catálogo", "Figura 2"),
    ("RF05", "Historial", "Lista ordenada por instante; la lectura nueva aparece tras el POST", "Figura 4 y sección 7.6"),
    ("RF06", "Alertas y anotaciones", "Alerta automática por umbral; página /anotaciones con rol funcional", "Figuras 5, 6 y 10 a 12"),
    ("RF07", "Modelo relacionado", "7 tablas con llaves foráneas, UNIQUE y CHECK", "Secciones 4 y 7.6"),
    ("RF08 / RF09", "API y IoT", "No aplican en Web 1.0; la columna origen ya distingue manual/simulado", "—"),
], anchos=[1.6, 3.2, 7.2, 4.5], tam=9)
p("De los criterios de aceptación de la ficha, en este incremento cubrimos CA01 (flujo completo con cambio "
  "persistido), CA02 (caso inválido rechazado sin alterar datos), CA05 (otra persona puede reproducirlo con el "
  "README) y CA06 (autoría y contribución). CA03 y CA04 corresponden a la API y al evento IoT de etapas posteriores.")

# 4
h("4. Modelo de datos")
p("La ficha propone seis entidades: zona, variable, lectura, umbral, alerta y anotación. Nosotros agregamos "
  "sensor, porque en la práctica una lectura no la produce una variable abstracta sino un dispositivo instalado "
  "en una zona, y ese dispositivo es justo lo que la ficha pide identificar cuando lleguemos a la simulación IoT "
  "(dispositivo, instante, variable, unidad y valor). Con eso quedamos en siete entidades, dentro del límite.")
figura("00_modelo_er.png", "Modelo de datos del PR09 en P02. Las flechas indican la dirección de la llave foránea.", 16)
p("Hay dos rangos distintos a propósito. El rango físico vive en variable y sirve para rechazar valores imposibles "
  "(150 °C en temperatura ambiente): eso es una validación y responde 400. El umbral vive en umbral, es por sensor, y "
  "sirve para aceptar la lectura pero avisar: 38 °C es un valor posible, se guarda, y como pasa del máximo operativo "
  "de 32 °C se genera una alerta ALTA. Esa separación nos permitió mostrar en el mismo incremento una validación "
  "negativa y una regla de negocio con efecto persistido.")
p("La columna origen de lectura solo admite 'manual' o 'simulado' (hay un CHECK en la base). Es nuestra respuesta "
  "al riesgo principal de la ficha: en la interfaz cada lectura lleva su etiqueta de procedencia y la unidad junto "
  "al valor, así que las dos filas de la semilla se ven claramente como simuladas y las que captura el observador "
  "como manuales. En anotación pasa algo parecido con autor_rol: la base solo acepta RESPONSABLE, OBSERVADOR o "
  "COORDINACION, que son los tres roles funcionales del proyecto.")
tabla(["Entidad", "Para qué sirve", "Restricciones que aplica PostgreSQL"], [
    ("zona", "Cama, invernadero o parcela del huerto", "nombre único; ubicación obligatoria"),
    ("variable", "Magnitud que se mide, con unidad y rango físico", "clave única; valor_minimo < valor_maximo"),
    ("sensor", "Dispositivo que mide una variable en una zona", "FK a zona y variable; código único"),
    ("umbral", "Rango operativo deseado de un sensor", "uno por sensor; minimo < maximo"),
    ("lectura", "Medición con instante, valor y procedencia", "FK a sensor; origen IN (manual, simulado); índice por sensor y fecha"),
    ("alerta", "Aviso generado por una lectura fuera de umbral", "una por lectura; nivel IN (BAJA, ALTA)"),
    ("anotacion", "Nota de un rol sobre una zona, opcionalmente ligada a una lectura", "FK a zona y lectura; rol IN (3 roles); texto de 3 a 300"),
], anchos=[2.6, 6.4, 7.5], tam=9)

# 5
h("5. Cómo está construida la aplicación")
p("Seguimos la separación que se explica en M02: la JSP solo presenta, el Servlet coordina la petición, una clase de "
  "servicio aplica las reglas y abre la transacción, y los repositorios concentran el SQL. Así pudimos probar las "
  "reglas de validación con JUnit sin levantar Tomcat ni la base de datos, y ninguna vista tiene SQL ni credenciales.")
tabla(["Capa", "Clases", "Responsabilidad"], [
    ("Configuración", "DbConfig, AppContextListener", "Lee DB_URL, DB_USER y DB_PASSWORD del entorno; libera el driver JDBC al parar"),
    ("Modelo", "Zona, Sensor, Lectura, Alerta, Anotacion", "Objetos simples que viajan del repositorio a la JSP"),
    ("Repositorio", "SensorRepository, LecturaRepository, AnotacionRepository", "PreparedStatement, try-with-resources, RETURNING id"),
    ("Servicio", "LecturaValidator, AnotacionValidator, LecturaService", "Reglas de RF03 y RF06; transacción lectura + alerta"),
    ("Web", "LecturaServlet, AlertaServlet, AnotacionServlet, HealthServlet", "GET/POST, códigos 200/303/400, forward a las vistas"),
    ("Vistas", "lecturas.jsp, alertas.jsp, anotaciones.jsp, error.jsp", "JSTL con c:out para escapar HTML"),
], anchos=[2.6, 6.2, 7.7], tam=9)
p("Un ejemplo de cómo insertamos la lectura y la alerta dentro de la misma conexión, con autocommit apagado:")
bloque_codigo(
    "try (Connection cn = DbConfig.getConnection()) {\n"
    "    cn.setAutoCommit(false);\n"
    "    try {\n"
    "        Sensor sensor = sensorRepository.findById(cn, sensorId).orElseThrow(...);\n"
    "        LecturaValidator.validarRangoFisico(sensor, valor, errores);   // 400 si falla\n"
    "        long lecturaId = lecturaRepository.insertLectura(cn, sensor.getId(), valor, observacion);\n"
    "        String nivel = LecturaValidator.nivelAlerta(sensor, valor);      // BAJA, ALTA o null\n"
    "        if (nivel != null) {\n"
    "            lecturaRepository.insertAlerta(cn, lecturaId, nivel, LecturaValidator.mensajeAlerta(sensor, valor, nivel));\n"
    "        }\n"
    "        cn.commit();\n"
    "    } catch (ValidacionException | SQLException | RuntimeException e) {\n"
    "        cn.rollback();\n"
    "        throw e;\n"
    "    }\n"
    "}")
p("Y el INSERT correspondiente, siempre parametrizado:")
bloque_codigo(
    "String sql = \"INSERT INTO lectura (sensor_id, valor, origen, observacion) VALUES (?, ?, 'manual', ?) RETURNING id\";\n"
    "try (PreparedStatement ps = cn.prepareStatement(sql)) {\n"
    "    ps.setLong(1, sensorId);\n"
    "    ps.setBigDecimal(2, valor);\n"
    "    ps.setString(3, observacion);\n"
    "    try (ResultSet rs = ps.executeQuery()) { rs.next(); return rs.getLong(1); }\n"
    "}")
tabla(["Ruta", "Método", "Qué hace"], [
    ("/web1/", "GET", "Redirige a /lecturas"),
    ("/web1/lecturas", "GET", "Catálogo de sensores con umbral e historial de lecturas"),
    ("/web1/lecturas", "POST", "Registra una lectura; 303 si es válida, 400 con errores si no"),
    ("/web1/alertas", "GET", "Alertas generadas"),
    ("/web1/anotaciones", "GET / POST", "Zonas y anotaciones; alta con rol funcional"),
    ("/web1/health", "GET", "JSON con estado de la conexión y versión de PostgreSQL; 503 si la base no responde"),
], anchos=[4, 2.5, 10], tam=9)

# 6
h("6. Instalación y despliegue")
p("El paso a paso completo, con las dos opciones de base de datos (Docker o PostgreSQL ya instalado), está en el "
  "README del repositorio. Aquí va el resumen de lo que hicimos nosotros en Windows con Git Bash:")
bloque_codigo(
    "git clone https://github.com/zywoxxx/dsw-p02-monitoreo-huerto.git\n"
    "cd dsw-p02-monitoreo-huerto\n"
    "docker compose -f docker/docker-compose.yml up -d        # PostgreSQL 16 en localhost:5436, crea tablas y semilla\n"
    "cp scripts/setenv.bat.example \"$CATALINA_HOME/bin/setenv.bat\"   # DB_URL, DB_USER, DB_PASSWORD fuera del código\n"
    "mvn clean package                                          # 14 pruebas, target/web1.war\n"
    "cp target/web1.war \"$CATALINA_HOME/webapps/\" && \"$CATALINA_HOME/bin/startup.bat\"\n"
    "curl http://localhost:8080/web1/health")
p("Las credenciales de laboratorio (usuario huerto_app, contraseña huerto_dev) están solo en el docker-compose y en "
  "los archivos de ejemplo de setenv; el setenv real de Tomcat está en .gitignore. La aplicación las lee de las "
  "variables de entorno en DbConfig, de modo que el mismo WAR sirve en cualquier máquina.")
p("Para no depender de que alguien repita todo a mano, hicimos scripts/verify-module.sh. Compila, corre las pruebas "
  "unitarias, revisa el contenido del WAR, consulta PostgreSQL, vuelve a desplegar en Tomcat esperando a que el "
  "contexto suba, y ejecuta con curl el recorrido completo (GET inicial, POST válido, tres POST inválidos, POST que "
  "genera alerta, GET con persistencia, alertas y anotaciones). Cada paso deja su salida en docs/evidencia/txt y se "
  "marca como VERIFICADO o NO_VERIFICADO. La última corrida terminó así:")
bloque_codigo(leer("16_resumen.txt"))

# 7
h("7. Pruebas y evidencia")
h("7.1 Pruebas unitarias", 2)
p("Las reglas de validación están en clases sin dependencias de Tomcat ni de JDBC, así que se prueban con JUnit 5 "
  "durante mvn package. Son 14 casos: 10 para lecturas (valor válido en rango, alerta ALTA, alerta BAJA, valor "
  "vacío, no numérico, tres decimales, fuera del rango físico, sensor inválido, sensor inactivo y observación "
  "demasiado larga) y 4 para anotaciones (caso válido, rol fuera de los tres permitidos, texto vacío o muy largo, "
  "zona y lectura inválidas).")
bloque_codigo(leer("02_pruebas_unitarias.txt"))

h("7.2 Recorrido del flujo principal", 2)
p("Las capturas se tomaron con un navegador real contra la aplicación desplegada, después de dejar la base con la "
  "semilla limpia. Primero la pantalla inicial: cinco sensores en dos zonas y las dos lecturas simuladas de la "
  "semilla, cada una con su etiqueta de procedencia.")
figura("01_get_inicial_lecturas.png", "GET /web1/lecturas con la semilla inicial (RF01, RF02, RF04, RF05).")
p("Después llenamos el formulario con el sensor de temperatura de la Cama A y 27.5 °C, un valor dentro del umbral.")
figura("02_formulario_lectura_valida.png", "Formulario antes de enviar (RF03).")
figura("03_post_valido_resultado.png", "Resultado del POST válido: mensaje de confirmación y la lectura nueva marcada como manual (CA01).")
p("Con 38 °C la lectura también es válida (está dentro del rango físico) pero supera el umbral de 32 °C, así que "
  "además del mensaje de éxito aparece el aviso de alerta y la fila queda resaltada.")
figura("07_post_valido_con_alerta.png", "POST con 38 °C: lectura guardada y alerta ALTA generada en la misma transacción (RF06).")
figura("08_get_alertas.png", "Vista /web1/alertas con la alerta recién creada.")
p("En la consola, la respuesta del servidor al POST válido fue la siguiente (el 303 y la cabecera Location son lo "
  "que demuestra el patrón POST-Redirect-GET):")
bloque_codigo(leer("08_post_valido.txt", max_lineas=5))

h("7.3 Validaciones negativas", 2)
p("Probamos los casos inválidos que marca la ficha y algunos más. En todos el servidor responde 400, muestra el "
  "motivo y conserva lo que ya se había escrito; en la base no cambia nada.")
figura("04_post_invalido_no_numerico.png", "Valor no numérico (abc): 400 con mensaje (CA02).")
figura("05_post_invalido_rango_fisico.png", "150 °C está fuera del rango físico de la variable: se rechaza (prueba mínima de la ficha).")
figura("06_post_invalido_vacio.png", "Campos vacíos: dos mensajes de error a la vez.")
p("Por curl también comprobamos un sensor inexistente (id 999), que responde 400 con 'El sensor seleccionado no "
  "existe'. Este es el encabezado de la respuesta al valor no numérico:")
bloque_codigo(leer("09_post_invalido_no_numerico.txt", max_lineas=4))

h("7.4 Zonas y anotaciones", 2)
p("La página de anotaciones lista las zonas del huerto y permite que cualquiera de los tres roles deje una nota "
  "sobre una zona, opcionalmente ligada al número de una lectura. El rol se valida contra la lista de tres y la "
  "base lo vuelve a comprobar con un CHECK.")
figura("11_get_zonas_anotaciones.png", "GET /web1/anotaciones con las zonas y la anotación de la semilla (RF01, RF06).")
figura("12_post_anotacion_valida.png", "Anotación registrada por el rol OBSERVADOR.")
figura("13_post_anotacion_invalida.png", "Texto de dos caracteres: 400 con el mensaje de longitud mínima.")

h("7.5 Despliegue en Tomcat y persistencia en PostgreSQL", 2)
p("La ruta de salud nos sirvió durante todo el desarrollo para saber si Tomcat estaba conectado a la base. Devuelve "
  "la versión del servidor PostgreSQL y el total de lecturas; si la base no responde, devuelve 503 con el motivo.")
p("Así se ve la respuesta con curl (la captura del navegador está en docs/evidencia/img/09_health_json.png, "
  "pero el JSON se lee mejor en texto):")
bloque_codigo(leer("06_health.txt"))
p("Del lado de la base de datos, estas son capturas de sesiones psql contra el PostgreSQL 16 del contenedor "
  "(puerto 5436) tomadas después de las capturas de la aplicación. En la primera se ve la conexión, las siete tablas "
  "con su dueño y los conteos; en la segunda, la estructura de lectura y anotación con sus llaves, restricciones "
  "CHECK e índices; en la tercera, las lecturas con su procedencia, la alerta y las anotaciones, que coinciden con "
  "las figuras anteriores.")
figura("14_psql_docker_tablas.png", "psql en PostgreSQL 16 (Docker): conexión, tablas y conteos por entidad.")
figura("17_psql_docker_estructura.png", "psql: estructura de las tablas lectura y anotación (llaves foráneas, CHECK e índices).")
figura("15_psql_docker_lecturas.png", "psql: lecturas con procedencia y alerta, alertas y anotaciones registradas.")
p("Estas son las líneas del log de Tomcat en las que se ve el despliegue del WAR (las rutas locales se sustituyeron "
  "por marcadores para no incluir rutas privadas):")
bloque_codigo(leer("05_tomcat_log.txt"))
p("También probamos la Opción B del README con el PostgreSQL 17.11 que uno de nosotros tiene instalado en Windows "
  "(puerto 5433): creamos el rol y la base con psql, corrimos los mismos scripts de esquema y semilla, cambiamos solo "
  "DB_URL en el setenv de Tomcat y la misma aplicación, sin recompilar, quedó conectada. Estas son las líneas "
  "clave de esa prueba (el archivo completo es docs/evidencia/txt/19_postgres_nativo_opcionB.txt):")
bloque_codigo("\n".join(l for l in leer("19_postgres_nativo_opcionB.txt").splitlines()
                          if l.startswith(("#", "$", " PostgreSQL 17", "HTTP/1.1", "Location", "  3 |", "  2 |", "  1 |", " id |", "----+", "(3 filas)"))
                          or "debe ser numerico" in l or l.startswith("{\"status\"")), tam=8)
figura("16_psql_nativo_pg17.png", "psql en el PostgreSQL 17 instalado en Windows (puerto 5433): las mismas siete tablas y la lectura registrada desde la aplicación.")
p("Como prueba negativa de infraestructura detuvimos el contenedor de PostgreSQL con la aplicación corriendo: "
  "/health pasó a 503 y /lecturas mostró la página de error controlada, sin trazas ni credenciales. Al volver a "
  "arrancar la base, la aplicación se recuperó sola sin reiniciar Tomcat.")
bloque_codigo(leer("18_health_sin_bd.txt", max_lineas=12))

# estado final de la base en una pagina horizontal, porque las filas de psql son anchas
from docx.enum.section import WD_SECTION
sh = doc.add_section(WD_SECTION.NEW_PAGE)
sh.orientation = WD_ORIENT.LANDSCAPE
sh.page_width, sh.page_height = Cm(27.94), Cm(21.59)
sh.left_margin = sh.right_margin = Cm(1.8)
sh.top_margin = sh.bottom_margin = Cm(2.0)
h("7.6 Estado de la base de datos después de las capturas", 2)
p("Consultado con psql usando sql/03_consultas_verificacion.sql. Se ven las siete tablas, los conteos, las lecturas "
  "con su procedencia, la alerta y las anotaciones. Las dos lecturas manuales y la anotación que empieza con "
  "'[prueba]' son exactamente las de las figuras anteriores; las dos lecturas simuladas son la semilla.")
bloque_codigo(leer("17_postgres_estado_final.txt", desde=5), tam=7.5)
sv = doc.add_section(WD_SECTION.NEW_PAGE)
sv.orientation = WD_ORIENT.PORTRAIT
sv.page_width, sv.page_height = Cm(21.59), Cm(27.94)
sv.left_margin = sv.right_margin = Cm(2.5)
sv.top_margin = sv.bottom_margin = Cm(2.2)

# 8
h("8. Problemas que tuvimos y cómo los resolvimos")
p("No todo salió a la primera y preferimos dejarlo registrado (la bitácora completa está en docs/bitacora.md).")
p("El primer tropiezo fue con las capturas automáticas. El script elegía el sensor por posición en la lista y el "
  "primero resultó ser el de humedad de suelo, así que un 27.5 que pensábamos que era una temperatura en rango salió "
  "con alerta BAJA. Lo corregimos seleccionando por el código del sensor y poniendo cada opción del select en una "
  "sola línea del HTML para poder leer su id desde los scripts.")
p("El segundo fue una condición de carrera con Tomcat. El script copiaba el WAR y consultaba la página de inmediato; "
  "Tomcat tarda unos segundos en replegar y volver a desplegar el contexto, y en ese lapso obtuvimos un 404 y hasta un "
  "500 que venía del propio Tomcat, no de nuestra aplicación. La solución fue esperar a que /health respondiera 404 "
  "después de retirar el WAR y 200 después de copiarlo. De paso agregamos un listener que anula el registro del "
  "driver JDBC al parar el contexto, porque Tomcat avisaba de una posible fuga de memoria.")
p("El tercero fue tonto pero ilustrativo: la comprobación de persistencia buscaba una marca en la salida de psql y la "
  "consulta no incluía la columna observación, así que fallaba aunque las filas sí estaban. Bastó agregar la columna "
  "a la consulta de verificación.")
p("El más importante fue de alcance. La primera versión del modelo tenía huerto y tipo_sensor y no tenía anotación; "
  "cuando revisamos la ficha PR09 con calma vimos que las entidades, la numeración de los requisitos y hasta la idea de "
  "'tres roles' (que son roles del sistema, no del equipo) no coincidían. Reescribimos el esquema, agregamos la "
  "página de anotaciones y volvimos a generar toda la evidencia antes de entregar.")

# 9
h("9. Relación con la rúbrica R02")
tabla(["Criterio R02", "Puntos", "Dónde está la evidencia"], [
    ("Navegación JSP/Servlet y flujo funcional", "2.5", "Figuras 2 a 12; respuestas HTTP en docs/evidencia/txt/07 a 14c; servlets y vistas en src/"),
    ("Despliegue correcto en Tomcat", "2", "Secciones 6 y 7.5; txt/01_build, 03_war_contenido, 05_tomcat_log, 06_health"),
    ("Persistencia en PostgreSQL", "2.5", "Secciones 4, 7.5 y 7.6; sql/; txt/04, 15, 17 y 18; repositorios con PreparedStatement"),
    ("README, capturas y repositorio", "2", "README.md, docs/ (acta, F02, trazabilidad, bitácora, reporte), docs/evidencia/INDICE.md, historial de commits"),
    ("Orden y limpieza técnica", "1", "Estructura Maven estándar, paquetes por capa, .gitignore sin target ni secretos, scripts de limpieza"),
], anchos=[5, 1.5, 10], tam=9)

# 10
h("10. Quién hizo qué")
p("Nos repartimos el trabajo en tres frentes y todos revisamos lo de los demás antes de cada commit.")
p("Juan Pablo Kuri Ricardez se encargó de la parte web: los servlets, las vistas JSP con JSTL, el empaquetado con "
  "Maven y el despliegue en Tomcat, incluida la configuración por variables de entorno.")
p("Pedro García Padilla llevó los datos: el modelo de siete entidades, los scripts de esquema, semilla, verificación "
  "y limpieza, y los repositorios JDBC con consultas parametrizadas y la transacción de lectura más alerta.")
p("Alejandro Pacheco Luna y Ariadna Trejo Álvarez se ocuparon de la evidencia y la documentación: el protocolo de "
  "pruebas HTTP y su automatización en verify-module.sh, las capturas con navegador, el README, la bitácora, la "
  "matriz de trazabilidad, el índice de evidencia frente a R02 y este documento. También revisaron la ficha PR09 "
  "contra el código, que fue donde salió el desajuste de la sección 8.")

# 11
h("11. Limitaciones y pendientes")
p("Declaramos por separado lo que sí verificamos y lo que no, como pide la guía. Todo lo de las secciones 6 y 7 "
  "quedó VERIFICADO en la máquina donde se desarrolló (Windows 11). Lo que queda como NO_VERIFICADO o PENDIENTE:")
tabla(["Punto", "Estado", "Comentario"], [
    ("Reproducción por otro integrante en una máquina limpia", "PENDIENTE", "Es el CA05; el README y el script están hechos para eso"),
    ("Despliegue en Linux o macOS con setenv.sh", "NO_VERIFICADO", "Solo probamos Windows con setenv.bat"),
    ("Tres decimales y sensor inactivo por HTTP", "NO_VERIFICADO", "Cubiertos por pruebas unitarias; la semilla no tiene sensores inactivos"),
    ("/health sin variables DB_* definidas", "NO_VERIFICADO", "Probamos base detenida, no configuración ausente"),
    ("Atender o desactivar alertas", "PENDIENTE", "Prueba mínima de la ficha que dejamos para el siguiente incremento"),
    ("Autenticación y autorización por rol", "PENDIENTE", "Fuera de Web 1.0; el rol se declara y se persiste, pero no se verifica identidad"),
], anchos=[6, 2.8, 7.7], tam=9)
p("Otras limitaciones conocidas: usamos una conexión JDBC por petición (sin pool), las fechas se muestran en la zona "
  "horaria del servidor de base de datos, y las zonas, sensores y umbrales se cargan por SQL porque su alta desde "
  "pantalla corresponde al responsable en el incremento de JSF.")

# anexo
h("Anexo. Comandos y archivos de referencia")
bloque_codigo(
    "# Verificación completa (desde la raíz del repositorio)\n"
    "export CATALINA_HOME=/ruta/a/apache-tomcat-9.0.115\n"
    "./scripts/verify-module.sh M02\n"
    "\n"
    "# Recorrido manual con curl\n"
    "B=http://localhost:8080/web1\n"
    "curl -i $B/lecturas\n"
    "curl -i -X POST -d \"sensorId=2&valor=25.5\" $B/lecturas\n"
    "curl -i -X POST -d \"sensorId=2&valor=abc\"  $B/lecturas\n"
    "curl -i -X POST -d \"sensorId=2&valor=38\"   $B/lecturas\n"
    "curl -i $B/alertas\n"
    "curl -i -X POST -d \"zonaId=1&autorRol=OBSERVADOR&texto=Hojas+con+manchas\" $B/anotaciones\n"
    "\n"
    "# Estado de la base y limpieza de datos de prueba\n"
    "docker exec -i dsw-p02-huerto-db psql -U huerto_app -d huerto_db < sql/03_consultas_verificacion.sql\n"
    "./scripts/cleanup.sh")
tabla(["Archivo", "Contenido"], [
    ("README.md", "Requisitos, instalación, despliegue, datos de prueba, verificación, decisiones de diseño"),
    ("docs/F02-seleccion-proyecto.md", "Comparación de PR09, PR03 y PR01 y plantilla publicada en F02"),
    ("docs/acta-proyecto.md", "Acta breve con roles funcionales, alcance, flujo y RF de la ficha"),
    ("docs/modelo-datos.md", "Diagrama ER y diccionario de datos"),
    ("docs/requisitos-trazabilidad.md", "RF → código → SQL → prueba → evidencia → R02"),
    ("docs/bitacora.md", "Fecha, comando, resultado, interpretación y siguiente acción de cada paso"),
    ("docs/verification-report.md", "Estado VERIFICADO / NO_VERIFICADO / PENDIENTE por prueba"),
    ("docs/evidencia/INDICE.md", "Cada captura y salida relacionada con un criterio de R02"),
], anchos=[6, 10.5], tam=9)

doc.save(OUT)
print("documento generado en", OUT)
