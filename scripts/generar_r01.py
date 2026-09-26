"""
generar_r01.py - Producto individual R01: comparativo Web 1.0-4.0, stacks de la EE y candidatos PRxx.
Genera docs/entrega/R01_PACHECO_ALEJANDRO.docx y la linea de tiempo docs/p03/evidencia/img/00_linea_tiempo_web.png.
Requiere python-docx y Pillow.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "entrega" / "R01_PACHECO_ALEJANDRO.docx"
LINEA = ROOT / "docs" / "p03" / "evidencia" / "img" / "00_linea_tiempo_web.png"
VERDE = RGBColor(0x2F, 0x6B, 0x3A)
GRIS = RGBColor(0x55, 0x55, 0x55)


# ---------- linea de tiempo (imagen) ----------
def dibujar_linea_tiempo():
    W, H = 1800, 540
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    def fuente(n, t):
        for f in (n, "arial.ttf"):
            try:
                return ImageFont.truetype(f, t)
            except OSError:
                pass
        return ImageFont.load_default()

    FT, FA, FS = fuente("arialbd.ttf", 26), fuente("arial.ttf", 19), fuente("arial.ttf", 17)
    verde, gris = (47, 107, 58), (90, 90, 90)
    d.line([(120, 250), (W - 120, 250)], fill=verde, width=5)
    etapas = [
        ("Web 1.0", "1991-2003", "Paginas estaticas y luego\ndinamicas en el servidor",
         "Usuario: lector. Formularios y\nrecarga completa. HTML, CGI, JSP/Servlets.", "En la EE: JSP + Tomcat (P02)"),
        ("Web 2.0", "2004-2010", "Participacion: el usuario\ntambien escribe",
         "Componentes con estado, AJAX y\nvalidacion en servidor. JSF/PrimeFaces.", "En la EE: JSF + PrimeFaces (P03)"),
        ("Web 3.0", "2011-2018", "Cliente separado de la API;\ndatos con significado",
         "SPA en el navegador, API REST con\nJSON, backend sin sesion. Angular/Spring.", "En la EE: Angular + Spring Boot"),
        ("Web 4.0", "2019-", "Contexto y dispositivos:\nla web reacciona a eventos",
         "Sensores y simuladores publican\neventos (MQTT); la app interpreta.", "En la EE: IoT simulado + PostgreSQL"),
    ]
    for i, (nombre, anios, resumen, detalle, ee) in enumerate(etapas):
        x = 200 + i * 420
        d.ellipse([x - 14, 236, x + 14, 264], fill=verde)
        d.text((x - 120, 196), nombre, font=FT, fill=verde)
        d.text((x - 120, 275), anios, font=FA, fill=gris)
        arriba = i % 2 == 0
        y = 28 if arriba else 322
        d.multiline_text((x - 120, y), resumen, font=FA, fill=(30, 30, 30), spacing=4)
        d.multiline_text((x - 120, y + 58), detalle, font=FS, fill=gris, spacing=3)
        d.text((x - 120, y + 110), ee, font=FS, fill=verde)
    d.text((80, H - 40), "Fechas orientativas segun Aghaei et al. (2012) y O'Reilly (2005); las etapas se traslapan.", font=FS, fill=gris)
    LINEA.parent.mkdir(parents=True, exist_ok=True)
    img.save(LINEA)


dibujar_linea_tiempo()

# ---------- documento ----------
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for nivel, tam in ((1, 16), (2, 13)):
    h_ = doc.styles[f"Heading {nivel}"]
    h_.font.name = "Calibri"
    h_.font.size = Pt(tam)
    h_.font.bold = True
    h_.font.color.rgb = VERDE
    h_.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.2)
hp = sec.header.paragraphs[0]
hp.text = "DSW-19559 · R01 · Alejandro Pacheco Luna · Evolución de la web y stacks de la EE"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRIS
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
        run._r.append(el)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = txt
        run._r.append(it)
run.font.size = Pt(9)


def p(texto, negrita=False, cursiva=False, tam=None, alinear=None, color=None, despues=None):
    par = doc.add_paragraph()
    r = par.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    if tam:
        r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color
    if alinear:
        par.alignment = alinear
    if despues is not None:
        par.paragraph_format.space_after = Pt(despues)
    return par


def h(texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def tabla(encabezados, filas, anchos=None, tam=9):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(e)
        r.bold = True
        r.font.size = Pt(tam)
        sombrear(c, "E6F2E8")
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(tam)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


for _ in range(5):
    doc.add_paragraph()
p("Universidad Veracruzana", negrita=True, tam=14, alinear=WD_ALIGN_PARAGRAPH.CENTER, color=VERDE, despues=0)
p("Desarrollo de Sistemas Web (DSW-19559)", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=30)
p("Evolución de la web (1.0 a 4.0) y stacks de la experiencia educativa", negrita=True, tam=20, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=6)
p("Comparativo, línea de tiempo y candidatos del banco de proyectos", tam=14, cursiva=True, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=40)
p("Producto individual (R01)", negrita=True, tam=13, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=4)
p("Alejandro Pacheco Luna", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p("Equipo DSW-E01", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
for _ in range(3):
    doc.add_paragraph()
p("Facilitador: Dr. Gabriel Rodríguez Vásquez", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER, despues=0)
p("Xalapa, Ver., 25 de septiembre de 2026", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER)
salto()

h("1. Para qué sirve este comparativo")
p("Antes de elegir un proyecto y de decidir con qué tecnología se construye cada incremento, conviene tener claro qué "
  "cambió entre una etapa de la web y la siguiente: quién hace qué (el servidor, el navegador, el usuario, los "
  "dispositivos), por dónde viajan los datos y qué problemas de seguridad aparecen. Este documento resume esa evolución, "
  "la relaciona con los cuatro stacks que usamos en la experiencia educativa y compara dos fichas del banco C11 como "
  "candidatas. La comparación no reserva ningún proyecto; la selección se confirma en el foro F02.")
p("Escribo desde la experiencia concreta de nuestro equipo: ya construimos el incremento Web 1.0 (JSP/Servlets) y el "
  "Web 2.0 (JSF/PrimeFaces) del PR09, así que varios ejemplos vienen de ahí.")

h("2. Matriz comparativa de las etapas")
tabla(["Dimensión", "Web 1.0", "Web 2.0", "Web 3.0", "Web 4.0"], [
    ("Propósito", "Publicar y consultar información; formularios básicos", "Participar: el usuario crea y modifica contenido; aplicaciones interactivas", "Separar cliente y servicios; datos con significado y reutilizables por máquinas", "Adaptarse al contexto: sensores, eventos y decisiones asistidas"),
    ("Arquitectura", "Servidor genera HTML completo por petición (CGI, JSP/Servlets)", "Servidor con componentes y estado de vista; actualizaciones parciales por AJAX", "Cliente SPA en el navegador + API REST/JSON sin estado + base de datos", "Fuentes de eventos (dispositivos o simuladores) + broker/colas + servicios que interpretan"),
    ("Interacción", "Recarga completa; un formulario, una respuesta", "Componentes ricos, validación en servidor con respuesta parcial", "Interfaz reactiva; el cliente decide qué pedir y cuándo", "La aplicación reacciona sin que el usuario pida: alertas, telemetría"),
    ("Papel del usuario", "Lector; llena formularios", "Autor y colaborador; sesión autenticada", "Usuario de una aplicación que consume varios servicios", "Participante en un entorno instrumentado; supervisa y corrige"),
    ("Flujo de datos", "Navegador -> servidor -> HTML", "Navegador <-> servidor (postback AJAX) -> base", "Navegador <-> API (JSON) -> base; varios clientes por API", "Dispositivo -> broker/API -> base -> vistas y alertas"),
    ("Seguridad típica", "Validar entrada, evitar SQL inyectado, sesión de contenedor", "ViewState contra CSRF, control de sesión, escape de salida", "Tokens (JWT/OAuth), CORS, validación en la API, HTTPS", "Autenticar dispositivos, etiquetar procedencia, no confundir simulación con medición real"),
    ("Tecnologías", "HTML, HTTP, CGI, PHP, JSP/Servlets", "AJAX, JSF/PrimeFaces, jQuery, XML/JSON", "Angular/React, REST, Spring Boot, JSON, OAuth", "MQTT, WebSockets, series de tiempo, reglas/ML"),
    ("En nuestra EE", "P02: JSP + Tomcat 9 + PostgreSQL", "P03: JSF 2.3 + PrimeFaces 12 + Tomcat 9", "Angular 16 + Spring Boot 2.7 + PostgreSQL", "Simulador IoT (MQTT) + PostgreSQL"),
], anchos=[2.5, 3.5, 3.5, 3.5, 3.5], tam=8.5)
p("Las fechas y los nombres de etapa son una convención (Aghaei, Nematbakhsh y Farsani, 2012; O'Reilly, 2005), no un "
  "estándar: una aplicación real mezcla rasgos de varias etapas. Nuestro PR09, por ejemplo, ya etiqueta la procedencia de "
  "cada lectura desde Web 1.0 pensando en la etapa 4.0.")

h("3. Línea de tiempo")
par = doc.add_paragraph()
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(str(LINEA), width=Cm(16.5))
c = doc.add_paragraph()
c.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = c.add_run("Figura 1. Evolución de la web y su correspondencia con los incrementos de la EE.")
r.italic = True
r.font.size = Pt(9.5)
r.font.color.rgb = GRIS
p("Qué cambia de una etapa a otra, en una frase cada una: de 1.0 a 2.0 el usuario pasa de leer a escribir y el "
  "servidor deja de reconstruir toda la página en cada clic; de 2.0 a 3.0 la interfaz se muda al navegador y el servidor "
  "se vuelve una API que otros clientes también pueden usar; de 3.0 a 4.0 la fuente de datos ya no es solo una persona "
  "frente a un formulario, sino dispositivos que publican eventos y una aplicación que los interpreta.")

h("4. Etapa, arquitectura y responsabilidades")
p("Web 1.0. El servidor hace todo: recibe la petición, consulta la base y devuelve HTML. En nuestro P02, el Servlet "
  "coordinaba y la JSP presentaba; la validación vivía en clases Java y respondía 400 con la lista de errores. El "
  "usuario solo ve páginas completas; cada envío es una recarga. Es simple de razonar y de probar con curl, pero cada "
  "interacción cuesta una página entera.")
p("Web 2.0. El servidor sigue mandando, pero mantiene el estado de la vista y responde a peticiones parciales. En P03 "
  "el bean de vista guarda lo capturado, PrimeFaces envía el formulario por AJAX y solo se actualizan mensajes, "
  "formulario y tabla. Aparecen responsabilidades nuevas: proteger el estado (ViewState), controlar la sesión con una "
  "persona autenticada y no con un rol elegido en un combo, y evitar que un componente redondee en silencio lo que el "
  "servidor debería rechazar. Ya no basta curl para probar: hay cookies, ViewState y respuestas parciales.")
p("Web 3.0. La interfaz se construye en el navegador (Angular) y el servidor expone una API (Spring Boot) que no guarda "
  "sesión: cada petición lleva su token. El usuario interactúa con una aplicación que se siente local; el servidor "
  "responde JSON con códigos 2xx y 4xx que el cliente debe mostrar. La misma API podría servir a una app móvil.")
p("Web 4.0. La aplicación recibe eventos de dispositivos (o de un simulador, en nuestro caso) con instante, variable, "
  "unidad y valor, los guarda y los interpreta con las mismas reglas de umbral que ya existen. El usuario supervisa y "
  "anota; la responsabilidad ética es no presentar una simulación como si fuera una medición real.")

h("5. Comparación de los stacks de la EE")
tabla(["Aspecto", "JSP / Servlets + Tomcat", "JSF 2.3 / PrimeFaces + Tomcat", "Angular + Spring Boot", "IoT simulado (MQTT) + PostgreSQL"], [
    ("Dónde se dibuja la interfaz", "Servidor (JSP)", "Servidor (Facelets), con actualización parcial", "Navegador (componentes Angular)", "Vistas de las etapas anteriores más paneles de eventos"),
    ("Estado", "Sesión HTTP y parámetros", "Beans @ViewScoped/@SessionScoped, ViewState", "Estado en el cliente; API sin estado", "Series de lecturas en la base; estado del dispositivo"),
    ("Validación", "Clases Java, respuesta 400", "Validadores y conversores JSF, mensajes por campo", "Cliente (formularios reactivos) y servidor (Bean Validation), 4xx", "Rango físico y procedencia al ingerir el evento"),
    ("Seguridad", "Sesión del contenedor, PreparedStatement", "ViewState contra CSRF, filtro de sesión, escape en Facelets", "JWT/OAuth, CORS, HTTPS", "Autenticación del publicador, etiquetado simulado/real"),
    ("Persistencia", "JDBC parametrizado", "JDBC parametrizado (reutilizado)", "JPA/JDBC en Spring Data", "Inserción por eventos; índices por sensor y fecha"),
    ("Cuándo conviene", "Aprender el ciclo HTTP; sistemas pequeños", "Formularios ricos con equipo Java; intranets", "Interfaces complejas y varios clientes", "Monitoreo, telemetría, alertas"),
    ("Coste que vimos", "Recarga completa; sin componentes", "Curva de CDI/scopes; jars extra (Weld, JAXB) en Tomcat 9", "Dos proyectos y dos despliegues; versiones de Node/Angular", "Necesita simulador y reglas claras de procedencia"),
], anchos=[2.6, 3.4, 3.6, 3.4, 3.5], tam=8.5)
p("Con PostgreSQL como base común, lo que se conserva entre etapas es el modelo y las reglas: en nuestro caso, las "
  "siete tablas, la transacción de lectura más alerta y los validadores pasaron de P02 a P03 sin cambios. Lo que cambia "
  "es la capa de presentación y la forma de asegurar la sesión.")

h("6. Dos candidatos del banco C11")
p("Comparo las dos fichas que más se parecen en patrón (lectura, umbral, alerta) y que el equipo consideró viables. La "
  "comparación no reserva el proyecto.")
tabla(["Criterio", "PR09 Monitoreo de huerto o ambiente", "PR03 Monitoreo de laboratorio de cómputo"], [
    ("Problema", "Las observaciones ambientales no se conservan con contexto ni permiten reconocer tendencias y alertas", "Incidencias y condiciones ambientales sin relación con equipos, horarios y mantenimiento"),
    ("Usuarios (3 roles)", "Responsable de huerto, observador, coordinación", "Encargado de laboratorio, soporte técnico, coordinación"),
    ("Flujo principal", "Elegir sensor -> capturar lectura -> validar -> guardar con instante y procedencia -> alerta si sale del umbral -> historial", "Registrar equipo/incidencia -> lectura ambiental -> historial y alerta"),
    ("Requisitos", "Zonas, variables, lecturas, umbrales, historial, alertas y anotaciones (RF01-RF07); API y IoT después", "Inventario, incidencias, estado, lecturas, historial y alertas; API y IoT después"),
    ("Entidades", "zona, variable, lectura, umbral, alerta, anotación (+ sensor)", "equipo, laboratorio, incidencia, lectura, alerta, usuario"),
    ("Viabilidad técnica", "Alta: variables físicas simples, reglas claras, simulación sin hardware", "Media: mezcla inventario e incidencias con telemetría; complejidad media-alta"),
    ("Riesgo", "Presentar simulación como medición real -> etiquetar procedencia", "Usar telemetría para vigilar personas -> limitar a condiciones del espacio"),
], anchos=[3, 6.8, 6.7], tam=8.5)
p("Recomendación razonada: PR09. Cubre el mismo patrón que PR03 con un dominio más acotado, su riesgo se controla con "
  "una decisión de diseño (la columna de procedencia) y cada etapa de la web le agrega una capacidad observable sin "
  "inventar otro sistema: captura (1.0), formularios y sesión (2.0), API y panel (3.0), sensores simulados (4.0). PR03 "
  "queda como segunda opción por compartir el patrón, aunque arrastra el inventario de equipos.")

h("7. Conclusiones propias")
p("Lo que más me quedó claro al pasar de P02 a P03 es que el cambio de etapa no es cambiar de librería: es mover "
  "responsabilidades. En Web 1.0 el servidor era responsable de todo y un curl bastaba para probar; en Web 2.0 el "
  "servidor guarda estado por vista y hay que probar con un navegador porque el formulario ya no es un POST plano. Lo "
  "que no se movió fue la base y las reglas, y eso es lo que hace posible que el proyecto sea acumulativo.")
p("También aprendí que la seguridad cambia de forma con cada etapa: en 1.0 me preocupaba el SQL inyectado, en 2.0 la "
  "sesión y el ViewState, y para 3.0 y 4.0 ya sé que vendrán tokens y la pregunta de si un dato viene de un sensor real "
  "o de un simulador. Elegir el stack no es elegir el más moderno, sino el que reparte las responsabilidades como el "
  "problema lo necesita.")

h("Bibliografía")
for ref in (
    "Aghaei, S., Nematbakhsh, M. A. y Farsani, H. K. (2012). Evolution of the World Wide Web: From Web 1.0 to Web 4.0. International Journal of Web & Semantic Technology, 3(1), 1-10. https://doi.org/10.5121/ijwest.2012.3101",
    "Berners-Lee, T., Cailliau, R., Luotonen, A., Nielsen, H. F. y Secret, A. (1994). The World-Wide Web. Communications of the ACM, 37(8), 76-82. https://doi.org/10.1145/179606.179671",
    "Berners-Lee, T., Hendler, J. y Lassila, O. (2001). The Semantic Web. Scientific American, 284(5), 34-43.",
    "O'Reilly, T. (2005). What Is Web 2.0: Design Patterns and Business Models for the Next Generation of Software. O'Reilly Media. https://www.oreilly.com/pub/a/web2/archive/what-is-web-20.html",
    "Fielding, R. T. (2000). Architectural Styles and the Design of Network-based Software Architectures (tesis doctoral). University of California, Irvine. https://ics.uci.edu/~fielding/pubs/dissertation/top.htm",
    "IETF (2022). RFC 9110: HTTP Semantics. https://www.rfc-editor.org/rfc/rfc9110",
    "Apache Software Foundation. Apache Tomcat 9 Documentation. https://tomcat.apache.org/tomcat-9.0-doc/",
    "Eclipse Foundation. Jakarta Servlet 4.0 Specification. https://jakarta.ee/specifications/servlet/4.0/",
    "Eclipse Foundation. Jakarta Server Faces 2.3 Specification. https://jakarta.ee/specifications/faces/2.3/",
    "PrimeTek. PrimeFaces 12 Documentation. https://primefaces.github.io/primefaces/12_0_0/",
    "Angular. Version compatibility. https://angular.dev/reference/versions",
    "VMware. Spring Boot 2.7.18 Reference Documentation. https://docs.spring.io/spring-boot/docs/2.7.18/reference/html/",
    "OASIS (2019). MQTT Version 5.0. https://docs.oasis-open.org/mqtt/mqtt/v5.0/mqtt-v5.0.html",
    "PostgreSQL Global Development Group. PostgreSQL 16 Documentation. https://www.postgresql.org/docs/16/",
    "OWASP Foundation. Application Security Verification Standard 5.0. https://owasp.org/www-project-application-security-verification-standard/",
    "W3C (2023). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
    "Rodríguez Vásquez, G. (2026). C11. Banco de proyectos integradores; M02 y M03, Desarrollo de Sistemas Web. Universidad Veracruzana (material del curso en Eminus).",
):
    par = doc.add_paragraph(ref)
    par.paragraph_format.left_indent = Cm(1)
    par.paragraph_format.first_line_indent = Cm(-1)
    par.paragraph_format.space_after = Pt(4)
    for r in par.runs:
        r.font.size = Pt(9.5)

doc.save(OUT)
print("documento generado en", OUT)
